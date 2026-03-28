import os
import json
from datetime import datetime
from pydantic import BaseModel, Field
from typing import Optional, Dict
import os, smtplib
import requests
from email.mime.text import MIMEText


from dotenv import load_dotenv
from fpdf import FPDF

import PyPDF2
from docx import Document

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_community.tools.tavily_search import TavilySearchResults
from langchain.tools import tool
from langchain.agents import create_agent

load_dotenv()

os.environ["GOOGLE_API_KEY"] = os.getenv("GOOGLE_API_KEY")
os.environ["TAVILY_API_KEY"] = os.getenv("TAVILY_API_KEY")
print("✓ API keys configured successfully!")

llm = ChatGoogleGenerativeAI(model = "gemini-2.5-flash")

file_name = "STEM Action.pdf"

def extract_text(file_name):
    if file_name.endswith(".pdf"):
        text = ""
        with open(file_name, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text

    elif file_name.endswith(".docx"):
        doc = Document(file_name)
        return "\n".join([p.text for p in doc.paragraphs])

    else:
        raise ValueError("Unsupported file format. Upload PDF or DOCX.")

raw_text = extract_text(file_name)
print("✓ Proposal text extracted")

import re

def extract_project_schema(proposal_text):
    prompt = f"""
    You are an expert nonprofit analyst.

    Extract structured project information from the proposal text below.
    The project may be in ANY sector (education, health, environment, arts, workforce, housing, etc.).

    If information is missing, leave it as an empty string.
    Do NOT hallucinate.

    Return ONLY valid JSON in this exact format:

    {{
        "organization_name": "",
        "project_name": "",
        "sector": "",
        "project_dates": "",
        "geographic_focus": "",
        "problem_statement_summary": "",
        "goals_objectives": "",
        "target_population": "",
        "key_activities": "",
        "expected_outcomes": "",
        "evaluation_methods": "",
        "organizational_capacity": "",
        "budget_summary": "",
        "funding_request": "",
        "sustainability_plan": "",
        "supporting_documents": ""
    }}

    Proposal Text:
    {proposal_text}
    """

    # Call the LLM
    response = llm.invoke(prompt)
    
    # Extract the JSON from response.content
    raw_text = response.content

    # Sometimes the model adds extra explanation, so we extract the JSON object
    import re
    match = re.search(r"\{.*\}", raw_text, re.DOTALL)
    if not match:
        raise ValueError(f"No JSON found in LLM response:\n{raw_text}")
    
    return json.loads(match.group())

project_schema = {
    "organization_name": "STEM Like A Girl",
    "project_name": "STEM IT UP",
    "project_dates": "May 2026 – June 2027",
    "venue": "Workshops at library, school, and institutions across Oregon, Washington, and New Jersey",
    "goals_objectives": "Expand access to high-quality, hands-on STEM learning by increasing confidence, curiosity, and engagement in STEM.",
    "target_audience": "Girls in grades 3–5 from underrepresented and low-income communities",
    "expected_impact": "At least a 25% increase in girls from underrepresented communities attending workshops.",
    "methodology": "Host 2 extra community-based workshops and provide partial to full financial assistances ",
    "evaluation_plan": "Surveys before and after the workshops",
    "team_background": "Board members worked in the tech field from Microsoft manager to students at Rutgers University",
    "budget_info": "Total project budget: $10,000 covering workshop supplies, trainings, financial assistances,and outreach.",
    "funding_request": "$5,000",
    "supporting_materials": "Positive testomonials from girls and parents."
}

proposal_schema =[
    "Executive Summary",
    "Problem Statement / Needs Assessment",
    "Project Goals & Objectives",
    "Methodology / Activities / Timeline",
    "Evaluation Plan & Success Metrics",
    "Team Background",
    "Budget & Funding Request",
    "Conclusion / Call to Action"
]

# Tool 1: Web Search Tool
search_tool = TavilySearchResults(
    max_results=5,
    search_depth="advanced",
    include_answer=True,
    name="web_search",
    time_range="year",
    description="Use this tool to search the web for active, real-world grant opportunities. Input should be a highly targeted search query."
)

# Tool 2: Custom Document Generation & PDF Output Tool
@tool
def generate_grant_and_save_pdf(funder_details: str) -> str:
    """
    Use this tool AFTER you have found a suitable grant from the web search.
    Pass the detailed funder information (Name, Mission, Priorities, Due Date, etc.) as a string to this tool.
    This tool will automatically generate the 8-section proposal using the LLM and save it directly as a PDF.
    """
    print("\n[Tool Executing] Generating Proposal Sections based on found grant...")
    full_proposal = {}
    
    # Generate each section utilizing the global LLM
    for section in proposal_schema:
        print(f" -> Generating: {section}")
        prompt = f"""
        You are an expert nonprofit grant writer. Write ONLY the section titled: {section}

        =============================
        FUNDER INFORMATION (from search)
        =============================
        {funder_details}

        =============================
        PROJECT INFORMATION
        =============================
        {json.dumps(project_schema, indent=2)}

        INSTRUCTIONS:
        - Align strongly with the funder's mission and funding priorities.
        - Use persuasive and impact-driven language.
        - 300-500 words.
        - Do NOT generate other section titles.
        """
        response = llm.invoke(prompt)
        full_proposal[section] = response.content

    print("\n[Tool Executing] Formatting and saving to PDF...")
    
    # Format Text
    date_str = datetime.now().strftime("%B %d, %Y")
    formatted_text = f"{project_schema['project_name']}\nGrant Proposal Submission\nDate: {date_str}\n\n{'='*70}\n"
    
    for title, content in full_proposal.items():
        formatted_text += f"\n\n{title.upper()}\n"
        formatted_text += "-" * len(title) + "\n\n"
        formatted_text += content.strip() + "\n"
        formatted_text += "\n" + "="*70 + "\n"

    # Save to PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", '', 11)

    # Clean text encoding to prevent latin-1 character errors in FPDF
    cleaned_text = formatted_text.encode('latin-1', 'replace').decode('latin-1')
    for line in cleaned_text.split("\n"):
        pdf.multi_cell(0, 8, line)

    pdf_file_name = "Grant_Proposal_Submission.pdf"
    pdf.output(pdf_file_name)
    
    return f"Success! The grant proposal has been generated and saved locally as {pdf_file_name}."

# Tool 3:Email text generator
from typing import Optional, Dict
from pydantic import BaseModel, Field
from langchain.tools import tool

class EmailDraftInput(BaseModel):
    org_name: str = Field(..., description="Organization/nonprofit name")
    goal: str = Field(..., description="What the email is trying to achieve")
    recipient_type: Optional[str] = Field("funder", description="Who this email is to (e.g., funder, partner, sponsor)")
    grant_name: Optional[str] = Field(None, description="Grant/funder name if relevant")
    context: Optional[str] = Field(None, description="Any extra context to include (optional)")

@tool("generate_email_draft", args_schema=EmailDraftInput)
def generate_email_draft(
    org_name: str,
    goal: str,
    recipient_type: Optional[str] = "funder",
    grant_name: Optional[str] = None,
    context: Optional[str] = None
) -> Dict[str, str]:
    """
    Generates an email draft for demo purposes.
    DOES NOT send email. Prints the draft.
    """

    prompt = f"""
Write a professional, friendly outreach email draft.

Output MUST be exactly:
SUBJECT: ...
BODY:
...

Rules:
- 120-–180 words
- Plain language
- Clear call-to-action
- Use placeholders like [Your Name], [Role], [Phone], [Website]

Organization: {org_name}
Recipient type: {recipient_type}
Goal: {goal}
Grant name (optional): {grant_name or "N/A"}
Extra context (optional): {context or "N/A"}
"""

    resp = llm.invoke(prompt).content.strip()

    subject = ""
    body = resp

    if "SUBJECT:" in resp:
        subject = resp.split("SUBJECT:", 1)[1].split("\n", 1)[0].strip()
    if "BODY:" in resp:
        body = resp.split("BODY:", 1)[1].strip()

    # print for demo
    print("\n================ EMAIL DRAFT ================")
    print("SUBJECT:", subject)
    print("\nBODY:\n" + body)
    print("============================================\n")

    return {"subject": subject, "body": body}

from datetime import datetime, timedelta
import os

from langchain_core.tools import tool

# Google Calendar API imports
from googleapiclient.discovery import build
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request

SCOPES = ["https://www.googleapis.com/auth/calendar.events"]

def _get_calendar_service():
    """
    Local dev: expects credentials.json (OAuth client) in repo, and will create token.json after you authorize once.
    Colab: same idea, but you upload credentials.json or mount Drive.
    """
    creds = None

    # token.json stores the user's access/refresh tokens after first auth
    if os.path.exists("token.json"):
        creds = Credentials.from_authorized_user_file("token.json", SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists("credentials.json"):
                raise FileNotFoundError(
                    "Missing credentials.json. Download OAuth client credentials from Google Calendar API quickstart."
                )
            flow = InstalledAppFlow.from_client_secrets_file("credentials.json", SCOPES)
            creds = flow.run_local_server(port=0)

        with open("token.json", "w") as token:
            token.write(creds.to_json())

    return build("calendar", "v3", credentials=creds)

@tool
def create_grant_deadline_event(
    deadline_date: str,
    title: str,
    application_url: str = "",
    timezone: str = "America/Los_Angeles",
) -> str:
    """
    Create an all-day Google Calendar event for a grant deadline.
    deadline_date format: YYYY-MM-DD
    """
    service = _get_calendar_service()

    # All-day event: use 'date' (not dateTime)
    start_date = deadline_date
    # end.date is exclusive for all-day events, so add 1 day
    end_date = (datetime.fromisoformat(deadline_date) + timedelta(days=1)).date().isoformat()

    description = f"Grant deadline.\n\nApply: {application_url}" if application_url else "Grant deadline."

    event = {
        "summary": title,
        "description": description,
        "start": {"date": start_date, "timeZone": timezone},
        "end": {"date": end_date, "timeZone": timezone},
    }

    created = service.events().insert(calendarId="primary", body=event).execute()
    return f"Created event: {created.get('htmlLink', '(no link returned)')}"

# Combine tools into a list for the agent
tools = [search_tool, generate_email_draft, generate_grant_and_save_pdf, create_grant_deadline_event]

from langgraph.prebuilt import create_react_agent

agent = create_react_agent(
    model=llm,
    tools=tools,
    prompt="""You are an autonomous expert grant and outreach assistant for nonprofits.

The current year is {current_year}. 

You have access to these tools:
- web_search: find relevant grants or funders. ALWAYS include "{current_year}" or "upcoming deadlines {current_year}" in your search queries to ensure you find active grants.
- generate_grant_and_save_pdf: create a proposal PDF
- generate_email_draft: create an outreach email draft
- create_grant_deadline_event: create a calendar event for a grant deadline

RULES:
- Be concise and professional
- Never hallucinate grant deadlines. If you cannot find a specific {current_year} deadline, state that clearly.
"""
)
